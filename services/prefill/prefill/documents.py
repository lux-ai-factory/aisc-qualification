"""An uploaded file, as text.

Four formats, because those are what people hand you: a PDF export, a Word
document, and plain text or markdown. A file that cannot be read raises, with
the reason in the message, so the form can say why instead of appearing to work.
"""
from __future__ import annotations

import io


class DocumentUnreadable(ValueError):
    """Nothing could be read out of this file, and guessing would be worse."""


def _decode(raw: bytes) -> str:
    """Text, whatever it was encoded as.

    A Word export saved as text is often cp1252; losing a curly quote is better
    than losing the document, so the last resort replaces what it cannot map.
    """
    for encoding in ("utf-8", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _from_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(raw))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as exc:
        raise DocumentUnreadable(f"this PDF could not be read: {exc}") from exc


def _from_docx(raw: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(raw))
    except Exception as exc:
        raise DocumentUnreadable(f"this .docx could not be read: {exc}") from exc
    # Headings come through as paragraphs, which is what the field reader wants:
    # a heading and the text under it, in the order they were written.
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


READERS = {
    "txt": _decode,
    "md": _decode,
    "markdown": _decode,
    "text": _decode,
    "pdf": _from_pdf,
    "docx": _from_docx,
}


def read_document(raw: bytes, filename: str) -> str:
    """The text of an uploaded file.

    The extension decides, because the bytes of a .docx and a .zip are the same
    bytes and refusing is better than opening whatever it turns out to be.
    """
    if not raw:
        raise DocumentUnreadable("the file is empty")
    extension = (filename or "").rsplit(".", 1)[-1].lower() if "." in (filename or "") else ""
    reader = READERS.get(extension)
    if reader is None:
        raise DocumentUnreadable(
            f"{extension or 'this file'} is not a format this reads: "
            f"{', '.join(sorted(set(READERS)))}"
        )
    text = reader(raw).strip()
    if not text:
        raise DocumentUnreadable("there is no text in this file")
    return text
