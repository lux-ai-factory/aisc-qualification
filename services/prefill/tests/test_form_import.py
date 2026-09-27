"""Reading a form's questions out of a file: .csv, .md or .docx.

Form-assembly spec (docs/superpowers/form-assembly-2026-09-24/01-spec.md),
R24 to R27. Text rules only, no model (R41).

Interface chosen here (the spec names parse_form_file):
  parse_form_file(raw: bytes, filename: str) -> result with
    .format     "csv" | "md" | "docx"
    .questions  list of items with .text, .citation, .required
    .warnings   list[str]
  A file it will not read raises prefill.documents.DocumentUnreadable (or a
  subclass), whose message the endpoint returns as the 422 detail.
"""
import io

import pytest

from prefill.documents import DocumentUnreadable

try:
    from prefill.form_import import parse_form_file
except ImportError as _missing:  # the spec's new module: until it exists each test fails, the suite runs
    def parse_form_file(*_args, _error=_missing, **_kwargs):
        raise _error


def rows(result):
    return [(q.text, q.citation, q.required) for q in result.questions]


ACME_CSV = (
    "question,citation,required\n"
    "Who signs off a model release?,Acme AI Policy §4.2,yes\n"
    '"Which datasets are approved, and by whom?",Acme AI Policy §5.1,\n'
    "How are incidents reported?,,no\n"
)


# ── R24 CSV ────────────────────────────────────────────────────────────────


class TestCsv:
    def test_r24_the_acme_example(self):
        result = parse_form_file(ACME_CSV.encode("utf-8"), "acme.csv")
        assert result.format == "csv"
        assert rows(result) == [
            ("Who signs off a model release?", "Acme AI Policy §4.2", True),
            ("Which datasets are approved, and by whom?", "Acme AI Policy §5.1", False),
            ("How are incidents reported?", "", False),
        ]
        assert result.warnings == []

    def test_r24_a_utf8_bom_is_not_part_of_the_header(self):
        result = parse_form_file(b"\xef\xbb\xbf" + ACME_CSV.encode("utf-8"), "acme.csv")
        assert len(result.questions) == 3
        assert result.questions[0].citation == "Acme AI Policy §4.2"

    @pytest.mark.parametrize("header", ["question", "Questions", " TEXT ", "Question text"])
    def test_r24_a_header_is_recognised_by_its_first_cell(self, header):
        result = parse_form_file(f"{header},citation\nWho signs off?,§4.2\n".encode(), "f.csv")
        assert rows(result) == [("Who signs off?", "§4.2", False)]

    def test_r24_header_columns_are_matched_by_name_in_any_order(self):
        text = "question,mandatory,reference\nWho signs off?,Y,§4.2\nWho audits?,,§9\n"
        assert rows(parse_form_file(text.encode(), "f.csv")) == [
            ("Who signs off?", "§4.2", True),
            ("Who audits?", "§9", False),
        ]

    @pytest.mark.parametrize("name", ["citation", "reference", "source", "clause"])
    def test_r24_every_citation_column_name(self, name):
        text = f"question,{name}\nWho signs off?,§4.2\n"
        assert rows(parse_form_file(text.encode(), "f.csv")) == [("Who signs off?", "§4.2", False)]

    def test_r24_without_a_header_the_columns_are_question_citation_required_and_extras_are_ignored(self):
        text = "Who signs off?,§4.2,true,extra,more\nWho audits?,§9,0,x\n"
        assert rows(parse_form_file(text.encode(), "f.csv")) == [
            ("Who signs off?", "§4.2", True),
            ("Who audits?", "§9", False),
        ]

    def test_r24_semicolons_when_the_first_line_has_them_and_no_comma(self):
        text = "Question;Citation;Required\nWho signs off?;Acme §4.2;required\nWho audits, and when?;§9;\n"
        assert rows(parse_form_file(text.encode(), "f.csv")) == [
            ("Who signs off?", "Acme §4.2", True),
            ("Who audits, and when?", "§9", False),
        ]

    def test_r24_a_semicolon_after_the_first_line_does_not_change_the_delimiter(self):
        text = "question,citation\nA; B?,§1\n"
        assert rows(parse_form_file(text.encode(), "f.csv")) == [("A; B?", "§1", False)]

    @pytest.mark.parametrize("value", ["yes", "Y", "TRUE", "1", "Required", " yes "])
    def test_r24_required_values(self, value):
        text = f"question,citation,required\nWho signs off?,,{value}\n"
        assert parse_form_file(text.encode(), "f.csv").questions[0].required is True

    @pytest.mark.parametrize("value", ["", "no", "n", "false", "0", "optional", "maybe"])
    def test_r24_a8_everything_else_is_not_required(self, value):
        text = f"question,citation,required\nWho signs off?,,{value}\n"
        assert parse_form_file(text.encode(), "f.csv").questions[0].required is False

    def test_r24_rows_with_a_blank_question_are_skipped_silently(self):
        text = "question,citation\n,§1\n   ,§2\nWho signs off?,§3\n\n"
        result = parse_form_file(text.encode(), "f.csv")
        assert rows(result) == [("Who signs off?", "§3", False)]
        assert result.warnings == []

    def test_r24_cp1252_is_read_too(self):
        text = "question,citation\nWho signs off the café’s model?,Acme §4.2\n"
        result = parse_form_file(text.encode("cp1252"), "f.csv")
        assert result.questions[0].text == "Who signs off the café’s model?"


# ── R25 Markdown ───────────────────────────────────────────────────────────

ACME_MD = """# Acme AI policy questionnaire

1. Who signs off a model release? [Acme AI Policy §4.2]
2. Which datasets are approved?
- How are incidents reported? [§7]
See [the policy](https://intranet/policy) for details.

| Question | Citation |
|---|---|
| Who may retrain the model? | Acme AI Policy §6.3 |
"""


class TestMarkdown:
    def test_r25_the_acme_example(self):
        result = parse_form_file(ACME_MD.encode("utf-8"), "acme.md")
        assert result.format == "md"
        assert rows(result) == [
            ("Who signs off a model release?", "Acme AI Policy §4.2", False),
            ("Which datasets are approved?", "", False),
            ("How are incidents reported?", "§7", False),
            ("See [the policy](https://intranet/policy) for details.", "", False),
            ("Who may retrain the model?", "Acme AI Policy §6.3", False),
        ]
        assert result.warnings == ["Skipped 1 heading."]

    def test_r25_markdown_extension_is_markdown_too(self):
        assert parse_form_file(b"- Who signs off?\n", "acme.markdown").format == "md"

    def test_r25_headings_are_counted_in_one_warning(self):
        text = "# Title\n## Part one\nWho signs off?\n### Part two\nWho audits?\n"
        result = parse_form_file(text.encode(), "f.md")
        assert [q.text for q in result.questions] == ["Who signs off?", "Who audits?"]
        assert result.warnings == ["Skipped 3 headings."]

    def test_r25_every_list_marker_and_bold_wrapper_is_removed(self):
        text = "- One?\n* Two?\n+ Three?\n1. Four?\n2) Five?\n**Six?**\n__Seven?__\n- **Eight?**\n"
        assert [q.text for q in parse_form_file(text.encode(), "f.md").questions] == [
            "One?", "Two?", "Three?", "Four?", "Five?", "Six?", "Seven?", "Eight?",
        ]

    def test_r25_code_fences_rules_and_comments_are_skipped(self):
        text = (
            "Who signs off?\n"
            "```\nnot a question\nnor this\n```\n"
            "---\n***\n"
            "<!-- a note to the editor -->\n"
            "Who audits?\n"
        )
        result = parse_form_file(text.encode(), "f.md")
        assert [q.text for q in result.questions] == ["Who signs off?", "Who audits?"]

    def test_r25_a_trailing_link_is_text_not_a_citation(self):
        text = "Read [the policy](https://intranet/policy)\n"
        assert rows(parse_form_file(text.encode(), "f.md")) == [
            ("Read [the policy](https://intranet/policy)", "", False)
        ]

    def test_r25_a_table_follows_the_csv_header_rules(self):
        text = (
            "| Question | Reference | Mandatory |\n"
            "|:---|---|---:|\n"
            "| Who signs off? | §4.2 | yes |\n"
            "| Who audits? | | |\n"
        )
        assert rows(parse_form_file(text.encode(), "f.md")) == [
            ("Who signs off?", "§4.2", True),
            ("Who audits?", "", False),
        ]


# ── R26 Word ───────────────────────────────────────────────────────────────


def acme_docx() -> bytes:
    """Built here with python-docx, as the existing prefill tests build theirs."""
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_paragraph("Acme questionnaire", style="Title")
    document.add_paragraph("Governance", style="Heading 1")
    document.add_paragraph("Who signs off a model release? [Acme AI Policy §4.2]", style="Normal")
    document.add_paragraph("Which datasets are approved?", style="List Paragraph")
    document.add_paragraph("")
    table = document.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ["Question", "Citation", "Required"]):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ["Who may retrain the model?", "§6.3", "yes"]):
        cell.text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


class TestWord:
    def test_r26_the_acme_example(self):
        result = parse_form_file(acme_docx(), "acme.docx")
        assert result.format == "docx"
        assert rows(result) == [
            ("Who signs off a model release?", "Acme AI Policy §4.2", False),
            ("Which datasets are approved?", "", False),
            ("Who may retrain the model?", "§6.3", True),
        ]
        assert result.warnings == ["Skipped 2 headings."]

    def test_r26_subtitles_are_skipped_too(self):
        docx = pytest.importorskip("docx")
        document = docx.Document()
        document.add_paragraph("A subtitle", style="Subtitle")
        document.add_paragraph("Who audits?")
        out = io.BytesIO()
        document.save(out)
        result = parse_form_file(out.getvalue(), "f.docx")
        assert [q.text for q in result.questions] == ["Who audits?"]
        assert result.warnings == ["Skipped 1 heading."]

    def test_r26_a_file_that_is_not_a_docx_is_refused(self):
        with pytest.raises(DocumentUnreadable):
            parse_form_file(b"PK\x03\x04 not really a zip", "f.docx")


# ── R27 limits and errors ──────────────────────────────────────────────────


class TestLimits:
    def test_r27_whitespace_in_text_and_citation_is_collapsed(self):
        text = 'question,citation\n"  Who   signs\toff?  ","  Acme   §4.2 "\n'
        assert rows(parse_form_file(text.encode(), "f.csv")) == [("Who signs off?", "Acme §4.2", False)]

    def test_r27_duplicates_keep_the_first_and_say_how_many(self):
        text = "question,citation\nWho signs off?,§1\nWHO  signs off?,§2\nWho audits?,§3\nwho audits?,§4\n"
        result = parse_form_file(text.encode(), "f.csv")
        assert rows(result) == [("Who signs off?", "§1", False), ("Who audits?", "§3", False)]
        assert "Removed 2 duplicate questions." in result.warnings

    def test_r27_one_duplicate_is_singular(self):
        result = parse_form_file(b"- Who?\n- who?\n", "f.md")
        assert "Removed 1 duplicate question." in result.warnings

    def test_r27_a_question_over_2000_characters_is_skipped_naming_its_line(self):
        long = "x" * 2001
        text = f"question,citation\nWho signs off?,§1\n{long},§2\n"
        result = parse_form_file(text.encode(), "f.csv")
        assert [q.text for q in result.questions] == ["Who signs off?"]
        assert "Question on line 3 is longer than 2000 characters and was skipped." in result.warnings

    def test_r27_a_csv_cell_over_the_csv_module_limit_is_skipped_not_an_error(self):
        # Verification F2: csv.reader's default field limit is 131072 characters;
        # a cell over it must follow the same rule as any question over 2000.
        huge = "z" * 200_000
        text = f"question,citation\nWho signs off?,§1\n{huge},§2\nWho audits?,§3\n"
        result = parse_form_file(text.encode(), "f.csv")
        assert [q.text for q in result.questions] == ["Who signs off?", "Who audits?"]
        assert "Question on line 3 is longer than 2000 characters and was skipped." in result.warnings

    def test_r27_a_huge_quoted_citation_is_cut_not_an_error(self):
        text = 'question,citation\nWho signs off?,"' + "c" * 200_000 + '"\n'
        result = parse_form_file(text.encode(), "f.csv")
        assert rows(result) == [("Who signs off?", "c" * 200, False)]
        assert "The citation on line 2 was shortened to 200 characters." in result.warnings

    def test_r27_exactly_2000_characters_is_kept(self):
        text = "question\n" + "y" * 2000 + "\n"
        assert len(parse_form_file(text.encode(), "f.csv").questions[0].text) == 2000

    def test_r27_a_citation_over_200_characters_is_cut_naming_its_line(self):
        text = "# Title\nWho signs off? [" + "c" * 250 + "]\n"
        result = parse_form_file(text.encode(), "f.md")
        assert result.questions[0].citation == "c" * 200
        assert "The citation on line 2 was shortened to 200 characters." in result.warnings

    def test_r27_more_than_200_questions_is_refused(self):
        text = "question\n" + "".join(f"Question {i}?\n" for i in range(201))
        with pytest.raises(DocumentUnreadable, match=r"^This file has more than 200 questions; split it into smaller forms\.$"):
            parse_form_file(text.encode(), "f.csv")

    def test_r27_exactly_200_is_fine(self):
        text = "question\n" + "".join(f"Question {i}?\n" for i in range(200))
        assert len(parse_form_file(text.encode(), "f.csv").questions) == 200

    @pytest.mark.parametrize("name,ext", [("policy.pdf", "pdf"), ("notes.txt", "txt"), ("book.xlsx", "xlsx")])
    def test_r27_a9_other_formats_are_refused(self, name, ext):
        with pytest.raises(DocumentUnreadable) as refused:
            parse_form_file(b"Who signs off?\n", name)
        assert str(refused.value) == f"{ext} is not a form format this reads: csv, docx, md"

    def test_r27_an_empty_file_is_refused(self):
        with pytest.raises(DocumentUnreadable, match="^the file is empty$"):
            parse_form_file(b"", "f.csv")

    def test_r27_a_file_with_no_questions_is_not_an_error(self):
        result = parse_form_file(b"# Only a heading\n", "f.md")
        assert result.questions == []
        assert result.warnings == ["Skipped 1 heading."]


# ── R41 no model ───────────────────────────────────────────────────────────


def test_r41_form_import_imports_only_the_standard_library_docx_and_this_package():
    import ast
    import sys
    from pathlib import Path

    source = Path(__file__).resolve().parents[1] / "prefill" / "form_import.py"
    assert source.exists(), source
    names = set()
    for node in ast.walk(ast.parse(source.read_text(encoding="utf-8"))):
        if isinstance(node, ast.Import):
            names |= {a.name.split(".")[0] for a in node.names}
        elif isinstance(node, ast.ImportFrom) and node.level == 0:
            names.add(node.module.split(".")[0])
    outside = {n for n in names if n not in sys.stdlib_module_names and n not in {"docx", "prefill", "__future__"}}
    assert outside == set()


# ── Addendum 06 (docs/superpowers/form-assembly-2026-09-24/06-spec-addendum.md) ──
# R53: the importer reads what the exporter writes (the Annex column, the CSV
# formula guard, Markdown escapes); R70: the .docx expansion cap.


def annexed(result):
    return [(q.text, q.citation, q.required, q.annex_point) for q in result.questions]


def annex_warning(line, value):
    return f'The Annex IV point on line {line} ("{value}") is not one of the 14 and was left blank.'


class TestAnnexColumn:
    @pytest.mark.parametrize("header", ["annex_point", "Annex point", " ANNEX IV POINT ", "annex"])
    def test_r53_every_annex_header_name_in_a_csv(self, header):
        text = f"question,citation,required,{header}\nWho signs off?,§4.2,yes,2a\n"
        assert annexed(parse_form_file(text.encode(), "f.csv")) == [("Who signs off?", "§4.2", True, "2a")]

    def test_r53_the_annex_column_is_matched_by_name_in_any_position(self):
        text = "question,annex,required\nWho signs off?,1de,no\n"
        assert annexed(parse_form_file(text.encode(), "f.csv")) == [("Who signs off?", "", False, "1de")]

    def test_r53_the_value_is_trimmed_and_lowercased_and_empty_is_none(self):
        text = "question,annex_point\nOne?, 2A \nTwo?,\nThree?,1GH\n"
        result = parse_form_file(text.encode(), "f.csv")
        assert [q.annex_point for q in result.questions] == ["2a", None, "1gh"]
        assert result.warnings == []

    def test_r53_a_value_that_is_not_one_of_the_14_is_blank_with_a_warning_naming_the_line(self):
        text = "question,annex_point\nOne?,3a\nTwo?,2a\n"
        result = parse_form_file(text.encode(), "f.csv")
        assert [q.annex_point for q in result.questions] == [None, "2a"]
        assert result.warnings == [annex_warning(2, "3a")]

    def test_r53_the_warning_quotes_the_trimmed_value(self):
        result = parse_form_file("question,annex\nOne?,  9z  \n".encode(), "f.csv")
        assert result.warnings == [annex_warning(2, "9z")]

    def test_r53_b2_a_citation_in_the_annex_column_is_not_an_id(self):
        result = parse_form_file('question,annex\nOne?,"annex iv(2)(a)"\n'.encode(), "f.csv")
        assert result.questions[0].annex_point is None
        assert result.warnings == [annex_warning(2, "annex iv(2)(a)")]

    def test_r53_only_kept_questions_are_checked(self):
        long = "x" * 2001
        text = f"question,annex\nSame?,2a\nsame?,3a\n{long},4b\n"
        result = parse_form_file(text.encode(), "f.csv")
        assert annexed(result) == [("Same?", "", False, "2a")]
        assert not any("Annex IV point" in w for w in result.warnings)

    def test_r53_the_annex_warning_is_in_line_order_with_the_other_per_line_warnings(self):
        cite = "c" * 201
        text = f"question,citation,annex\nOne?,,3a\nTwo?,{cite},\nThree?,,7q\n"
        result = parse_form_file(text.encode(), "f.csv")
        assert result.warnings == [
            annex_warning(2, "3a"),
            "The citation on line 3 was shortened to 200 characters.",
            annex_warning(4, "7q"),
        ]

    def test_r53_without_a_header_row_there_is_no_annex_column(self):
        result = parse_form_file("Who signs off?,§4.2,yes,2a\n".encode(), "f.csv")
        assert annexed(result) == [("Who signs off?", "§4.2", True, None)]
        assert result.warnings == []

    def test_r53_no_is_not_required(self):
        result = parse_form_file("question,required\nOne?,no\n".encode(), "f.csv")
        assert result.questions[0].required is False

    def test_r53_a_markdown_table_header_names_the_annex_column(self):
        text = "| Question | Citation | Required | Annex IV point |\n|---|---|---|---|\n| Who signs off? | §4.2 | yes | 2g |\n"
        assert annexed(parse_form_file(text.encode(), "f.md")) == [("Who signs off?", "§4.2", True, "2g")]

    def test_r53_markdown_line_numbers_in_the_annex_warning(self):
        text = "# Title\n\n| Question | Annex |\n|---|---|\n| One? | 2a |\n| Two? | 3z |\n"
        result = parse_form_file(text.encode(), "f.md")
        assert [q.annex_point for q in result.questions] == ["2a", None]
        assert result.warnings == ["Skipped 1 heading.", annex_warning(6, "3z")]

    def test_r53_paragraph_and_list_questions_have_no_annex_point(self):
        result = parse_form_file(b"- Who signs off? [2a]\nWho audits?\n", "f.md")
        assert [q.annex_point for q in result.questions] == [None, None]

    def test_r53_a_docx_table_header_names_the_annex_column(self):
        docx = pytest.importorskip("docx")
        document = docx.Document()
        table = document.add_table(rows=2, cols=4)
        for cell, value in zip(table.rows[0].cells, ["Question", "Citation", "Required", "Annex"]):
            cell.text = value
        for cell, value in zip(table.rows[1].cells, ["Who may retrain the model?", "§6.3", "no", "2F"]):
            cell.text = value
        out = io.BytesIO()
        document.save(out)
        assert annexed(parse_form_file(out.getvalue(), "f.docx")) == [("Who may retrain the model?", "§6.3", False, "2f")]

    def test_r53_imported_question_has_annex_point(self):
        from prefill.form_import import ImportedQuestion

        assert "annex_point" in {f for f in ImportedQuestion.__dataclass_fields__}


class TestCsvFormulaGuard:
    @pytest.mark.parametrize("first", ["=", "+", "-", "@"])
    def test_r53_a_leading_quote_before_a_formula_character_is_removed_in_text_and_citation(self, first):
        text = f"question,citation\n'{first}x?,'{first}cite\n"
        assert rows(parse_form_file(text.encode(), "f.csv")) == [(f"{first}x?", f"{first}cite", False)]

    def test_r53_only_one_quote_is_removed(self):
        assert rows(parse_form_file("question\n''=x?\n".encode(), "f.csv")) == [("'=x?", "", False)]

    def test_r53_a_quote_before_anything_else_stays(self):
        assert rows(parse_form_file("question,citation\n'quoted' word?,'plain\n".encode(), "f.csv")) == [
            ("'quoted' word?", "'plain", False)
        ]

    def test_r53_markdown_cells_are_not_touched_by_the_guard(self):
        text = "| Question | Citation |\n|---|---|\n| '=x? | '+cite |\n"
        assert rows(parse_form_file(text.encode(), "f.md")) == [("'=x?", "'+cite", False)]

    def test_r53_docx_cells_are_not_touched_by_the_guard(self):
        docx = pytest.importorskip("docx")
        document = docx.Document()
        table = document.add_table(rows=2, cols=1)
        table.rows[0].cells[0].text = "Question"
        table.rows[1].cells[0].text = "'=x?"
        out = io.BytesIO()
        document.save(out)
        assert rows(parse_form_file(out.getvalue(), "f.docx")) == [("'=x?", "", False)]


class TestMarkdownEscapes:
    def test_r53_an_escaped_pipe_is_a_pipe_not_a_cell_boundary(self):
        text = "| Question | Citation |\n|---|---|\n| Who signs off \\| approves? | §4 |\n"
        assert rows(parse_form_file(text.encode(), "f.md")) == [("Who signs off | approves?", "§4", False)]

    def test_r53_a_double_backslash_is_one_backslash(self):
        text = "| Question | Citation |\n|---|---|\n| a\\\\\\|b | c\\\\d |\n"
        assert rows(parse_form_file(text.encode(), "f.md")) == [("a\\|b", "c\\d", False)]

    def test_r53_any_other_backslash_is_kept(self):
        text = "| Question | Citation |\n|---|---|\n| a\\nb \\x? | c:\\path |\n"
        assert rows(parse_form_file(text.encode(), "f.md")) == [("a\\nb \\x?", "c:\\path", False)]

    def test_r53_a_backslash_before_the_closing_pipe_of_a_cell(self):
        text = "| Question | Citation |\n|---|---|\n| ends in a backslash\\\\ | x |\n"
        assert rows(parse_form_file(text.encode(), "f.md")) == [("ends in a backslash\\", "x", False)]


# ── R70 .docx expansion cap ────────────────────────────────────────────────


def docx_bomb(total_mib: int = 51) -> bytes:
    """A real .docx plus one highly compressible member, declaring total_mib MiB unpacked."""
    import zipfile

    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_paragraph("Who signs off a model release?")
    base = io.BytesIO()
    document.save(base)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base.getvalue())) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("word/media/zeros.bin", b"\0" * (total_mib * 1024 * 1024))
    return out.getvalue()


def small_docx() -> bytes:
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_paragraph("Who audits?")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


EXPANDS = "this .docx expands to more than {} bytes when unpacked; remove embedded media or split it"


class TestDocxExpansion:
    def test_r70_a_docx_that_unpacks_to_more_than_50_mib_is_refused_before_docx_reads_it(self, monkeypatch):
        import docx

        def never(*_a, **_k):
            raise AssertionError("docx.Document was called")

        monkeypatch.delenv("PREFILL_MAX_UNZIPPED_BYTES", raising=False)
        bomb = docx_bomb()
        monkeypatch.setattr(docx, "Document", never)
        with pytest.raises(DocumentUnreadable) as refused:
            parse_form_file(bomb, "big.docx")
        assert str(refused.value) == EXPANDS.format(52428800)

    def test_r70_a_normal_docx_is_read_as_before(self, monkeypatch):
        monkeypatch.delenv("PREFILL_MAX_UNZIPPED_BYTES", raising=False)
        assert [q.text for q in parse_form_file(small_docx(), "f.docx").questions] == ["Who audits?"]

    def test_r70_the_limit_is_read_from_the_environment_at_call_time(self, monkeypatch):
        monkeypatch.setenv("PREFILL_MAX_UNZIPPED_BYTES", "1000")
        with pytest.raises(DocumentUnreadable) as refused:
            parse_form_file(small_docx(), "f.docx")
        assert str(refused.value) == EXPANDS.format(1000)

    def test_r70_bytes_that_are_not_a_zip_are_unreadable_with_the_existing_message(self):
        with pytest.raises(DocumentUnreadable) as refused:
            parse_form_file(b"PK\x03\x04 not really a zip", "f.docx")
        assert str(refused.value).startswith("this .docx could not be read: ")
